import tkinter as tk
from tkinter import filedialog, messagebox
from datetime import date

import openpyxl


students = []
attendance = []
current_index = 0
selected_roster_path = None


def choose_roster():
    global students, attendance, current_index, selected_roster_path

    file_path = filedialog.askopenfilename(
        title="Choose a class roster",
        filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")],
    )

    if not file_path:
        return

    try:
        selected_roster_path = file_path
        workbook = openpyxl.load_workbook(
            file_path,
            read_only=True,
            data_only=True,
            keep_vba=file_path.lower().endswith(".xlsm"),
        )
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        workbook.close()

        students = []
        for row in rows[1:]:
            first_name = row[0] if len(row) > 0 else None
            last_name = row[1] if len(row) > 1 else None
            if first_name or last_name:
                students.append((str(first_name or "").strip(), str(last_name or "").strip()))

        if not students:
            raise ValueError("No students were found below the header row.")

        attendance = [None] * len(students)
        current_index = 0
        preview_label.config(text=f"Loaded {len(students)} students from {worksheet.title}.")
        show_current_student()
    except Exception as error:
        messagebox.showerror("Could not open roster", str(error))


def record_attendance(value):
    if not students:
        return

    global current_index
    attendance[current_index] = value
    current_index += 1

    if current_index >= len(students):
        present = attendance.count(1)
        absent = attendance.count(0)
        other = attendance.count(-1)
        student_label.config(text="Finished!")
        progress_label.config(text=f"Present: {present}   Absent: {absent}   Not following: {other}")
        save_button.pack(pady=3)
        return

    show_current_student()


def show_current_student():
    first_name, last_name = students[current_index]
    student_label.config(text=f"{first_name} {last_name}".strip())
    progress_label.config(text=f"Student {current_index + 1} of {len(students)}")


def go_previous():
    global current_index

    if not students or current_index == 0:
        messagebox.showinfo("Already at first student", "There is no previous student.")
        return

    current_index -= 1
    attendance[current_index] = None
    save_button.pack_forget()
    show_current_student()


def save_attendance():
    if not selected_roster_path:
        messagebox.showwarning("No roster selected", "Choose a roster first.")
        return

    if not students or any(status is None for status in attendance):
        messagebox.showwarning("Incomplete attendance", "Please finish marking every student first.")
        return

    attendance_date = date.today().isoformat()

    workbook = openpyxl.load_workbook(
        selected_roster_path,
        keep_vba=selected_roster_path.lower().endswith(".xlsm"),
    )
    worksheet = workbook.active

    headers = [cell.value for cell in worksheet[1]]
    if attendance_date in headers:
        date_column = headers.index(attendance_date) + 1
    else:
        date_column = worksheet.max_column + 1
        worksheet.cell(row=1, column=date_column, value=attendance_date)

    existing_students = {}
    for row_number in range(2, worksheet.max_row + 1):
        first_name = worksheet.cell(row=row_number, column=1).value
        last_name = worksheet.cell(row=row_number, column=2).value
        existing_students[(str(first_name or "").strip(), str(last_name or "").strip())] = row_number

    for (first_name, last_name), status in zip(students, attendance):
        row_number = existing_students.get((first_name, last_name))
        if row_number is None:
            row_number = worksheet.max_row + 1
            worksheet.cell(row=row_number, column=1, value=first_name)
            worksheet.cell(row=row_number, column=2, value=last_name)
        worksheet.cell(row=row_number, column=date_column, value=status)

    workbook.save(selected_roster_path)
    messagebox.showinfo("Attendance saved", f"Saved attendance to:\n{selected_roster_path}")


def handle_key(event):
    if event.char == "1":
        record_attendance(1)
    elif event.char == "0":
        record_attendance(0)
    elif event.char == "-":
        record_attendance(-1)


window = tk.Tk()
window.title("Attendance App")
window.geometry("700x500")
window.bind("<Key>", handle_key)

header_frame = tk.Frame(window)
header_frame.pack(fill="x", padx=25, pady=(25, 10))

title_label = tk.Label(
    header_frame,
    text="Attendance App",
    font=("Arial", 24, "bold"),
)
title_label.grid(row=0, column=0, padx=20)

header_frame.grid_columnconfigure(0, weight=1)

preview_label = tk.Label(
    window,
    text="",
    font=("Consolas", 11),
    justify="left",
    anchor="w",
)
preview_label.pack(pady=5)

student_label = tk.Label(window, text="", font=("Arial", 65, "bold"))
student_label.pack(expand=True, pady=30)

bottom_frame = tk.Frame(window)
bottom_frame.pack(side="bottom", pady=20)

previous_button = tk.Button(
    bottom_frame,
    text="Previous",
    width=14,
    command=go_previous,
)
previous_button.grid(row=0, column=0, padx=15)

roster_controls = tk.Frame(bottom_frame)
roster_controls.grid(row=0, column=1, padx=15)

choose_button = tk.Button(
    header_frame,
    text="Choose Roster",
    font=("Arial", 12),
    command=choose_roster,
)
choose_button.grid(row=0, column=1, padx=20)

progress_label = tk.Label(
    roster_controls,
    text="",
    font=("Arial", 12),
)
progress_label.pack(pady=3)

save_button = tk.Button(
    roster_controls,
    text="Save Attendance to Roster",
    font=("Arial", 12),
    command=save_attendance,
)

window.mainloop()
